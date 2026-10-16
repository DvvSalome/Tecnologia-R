#!/usr/bin/env python3
"""Lee el Excel de productos, descarga imágenes, quita fondo (rembg) y genera JSON intermedios."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from pathlib import Path
from urllib.parse import unquote, urlparse
from urllib.request import Request, urlopen

import openpyxl
from PIL import Image
from rembg import remove

ROOT = Path(__file__).resolve().parents[1]
EXCEL_CANDIDATES = [
    ROOT / "data-excel.xlsx",
    Path.home() / "Downloads" / "Categoria de productos.xlsx",
]
DATA = ROOT / "data"
RAW_DIR = DATA / "images-raw"
NOBG_DIR = DATA / "images-nobg"
RAW_JSON = DATA / "excel-raw.json"
NORM_JSON = DATA / "products-normalized.json"
MANIFEST = DATA / "images-manifest.json"

CATEGORY_TO_TIPO = {
    "diademas": 4,
    "auriculares": 4,
    "audifonos": 4,
    "microfonos": 6,
    "teclados": 1,
    "mouses": 2,
    "monitores": 3,
    "camaras": 5,
    "convertidores": 10,
    "convertidores de video": 10,
    "bases": 11,
    "accesorios": 11,
    "reguladores": 12,
    "energia": 12,
    "tarjetas de sonido": 13,
    "audio": 13,
}


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text or "producto"


def fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.lower().strip()


def tipo_from_category(cat: str | None) -> int:
    if not cat:
        return 9
    f = fold(cat)
    # "Categoría X/Y" → toma la parte después de /
    if "/" in f:
        f = f.split("/", 1)[1].strip()
    f = f.replace("categoria ", "").strip()
    for key, tipo in CATEGORY_TO_TIPO.items():
        if key in f:
            return tipo
    if "convertidor" in f:
        return 10
    if "regulador" in f:
        return 12
    if "microfono" in f:
        return 6
    if "teclado" in f:
        return 1
    return 9


def parse_product_cell(cell: str) -> tuple[str, str]:
    """Separa nombre y descripción desde la celda Caracteristicas."""
    text = (cell or "").replace("\r\n", "\n").strip()
    if not text:
        return "Producto", ""
    # Quitar doble pipes del dump y normalizar
    lines = [ln.strip() for ln in text.split("\n")]
    lines = [ln for ln in lines if ln]
    if not lines:
        return "Producto", ""

    nombre = lines[0]
    # A veces la fila 0 es solo el nombre; a veces "Nombre | Características"
    if "Características" in nombre or "Caracteristicas" in nombre:
        parts = re.split(r"\s*\|\s*", nombre)
        nombre = parts[0].strip() or "Producto"
        rest = "\n".join(parts[1:] + lines[1:])
    else:
        rest = "\n".join(lines[1:])

    # Limpiar restos de "Características:"
    rest = re.sub(r"(?i)^\s*características\s*:?\s*", "", rest).strip()
    rest = rest.replace(" | ", "\n")
    # Compactar líneas vacías
    desc_lines = [ln.strip(" |") for ln in rest.split("\n")]
    desc_lines = [ln for ln in desc_lines if ln and ln.lower() not in {"características:", "caracteristicas:"}]
    descripcion = "\n".join(desc_lines).strip()
    if descripcion and not descripcion.lower().startswith("características"):
        descripcion = "Características:\n\n" + descripcion
    if "Garantía" not in descripcion and "garantia" not in fold(descripcion):
        descripcion = (descripcion + "\n\nGarantía 06 meses").strip()
    return nombre.strip(), descripcion


def download(url: str, dest: Path) -> Path | None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    try:
        req = Request(
            url,
            headers={
                "User-Agent": "Mozilla/5.0 (compatible; TecnologiaRCatalog/1.0)",
                "Accept": "image/*,*/*",
            },
        )
        with urlopen(req, timeout=45) as resp:
            data = resp.read()
        dest.write_bytes(data)
        return dest
    except Exception as e:
        print(f"  ⚠️  download fail {url}: {e}")
        return None


def remove_bg(src: Path, dest: Path) -> Path | None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    try:
        raw = Image.open(src).convert("RGBA")
        out = remove(raw)
        if not isinstance(out, Image.Image):
            out = Image.open(out).convert("RGBA")
        else:
            out = out.convert("RGBA")
        # Recortar transparentes y limitar lado máximo
        bbox = out.getbbox()
        if bbox:
            out = out.crop(bbox)
        max_side = 1200
        w, h = out.size
        scale = min(1.0, max_side / max(w, h))
        if scale < 1:
            out = out.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
        # Canvas cuadrado transparente
        side = max(out.size)
        canvas = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        canvas.paste(out, ((side - out.size[0]) // 2, (side - out.size[1]) // 2), out)
        canvas.save(dest, "PNG", optimize=True)
        return dest
    except Exception as e:
        print(f"  ⚠️  rembg fail {src.name}: {e}")
        return None


def find_excel() -> Path:
    for p in EXCEL_CANDIDATES:
        if p.exists():
            return p
    raise SystemExit("No se encontró el Excel (data-excel.xlsx o Downloads/Categoria de productos.xlsx)")


def main() -> None:
    excel = find_excel()
    print(f"📘 Excel: {excel}")
    DATA.mkdir(parents=True, exist_ok=True)
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    NOBG_DIR.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(excel, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows_raw = []
    products: dict[str, dict] = {}
    current_category = None

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        cat, features, valor, _imgref, imgurl = (list(row) + [None] * 5)[:5]
        if cat and str(cat).strip():
            current_category = str(cat).strip()
        if not features or not str(features).strip():
            continue
        if not imgurl or not str(imgurl).startswith("http"):
            continue

        nombre, descripcion = parse_product_cell(str(features))
        slug = slugify(nombre)
        tipo = tipo_from_category(current_category)
        precio = 0.0
        if valor not in (None, "", "#VALUE!"):
            try:
                precio = float(str(valor).replace(",", "").replace("$", "").strip())
            except ValueError:
                precio = 0.0

        rows_raw.append(
            {
                "row": i,
                "categoria": current_category,
                "nombre": nombre,
                "descripcion": descripcion,
                "precio": precio,
                "imageUrl": str(imgurl).strip(),
                "tipo": tipo,
                "slug": slug,
            }
        )

        p = products.get(slug)
        if not p:
            p = {
                "slug": slug,
                "nombre": nombre,
                "descripcion": descripcion,
                "tipo": tipo,
                "categoriaExcel": current_category,
                "precio": precio,
                "descuento": 0,
                "popular": False,
                "imageUrls": [],
            }
            products[slug] = p
        if str(imgurl).strip() not in p["imageUrls"]:
            p["imageUrls"].append(str(imgurl).strip())

    # Marcar populares: primer producto de cada tipo
    seen_tipos: set[int] = set()
    for p in products.values():
        if p["tipo"] not in seen_tipos:
            p["popular"] = True
            seen_tipos.add(p["tipo"])

    RAW_JSON.write_text(json.dumps({"rows": rows_raw}, ensure_ascii=False, indent=2), encoding="utf-8")
    NORM_JSON.write_text(json.dumps(list(products.values()), ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"✅ {len(products)} productos normalizados")

    manifest = []
    for p in products.values():
        local_images = []
        for idx, url in enumerate(p["imageUrls"]):
            ext = Path(unquote(urlparse(url).path)).suffix.lower() or ".img"
            if ext not in {".png", ".jpg", ".jpeg", ".webp", ".avif", ".gif"}:
                ext = ".img"
            raw_name = f"{p['slug']}-{idx}{ext}"
            raw_path = RAW_DIR / raw_name
            print(f"⬇️  {p['nombre']} [{idx}]")
            if not download(url, raw_path):
                continue
            nobg_path = NOBG_DIR / f"{p['slug']}-{idx}.png"
            if remove_bg(raw_path, nobg_path):
                local_images.append(str(nobg_path.relative_to(ROOT)))
        manifest.append(
            {
                "slug": p["slug"],
                "nombre": p["nombre"],
                "tipo": p["tipo"],
                "descripcion": p["descripcion"],
                "precio": p["precio"],
                "descuento": p["descuento"],
                "popular": p["popular"],
                "localImages": local_images,
                "sourceUrls": p["imageUrls"],
            }
        )

    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"✅ Manifest: {MANIFEST} ({len(manifest)} items)")


if __name__ == "__main__":
    main()
